import os
import json
import csv
import gzip
import sqlite3
import logging
from concurrent.futures import ProcessPoolExecutor, as_completed
from typing import Dict, List, Any, Optional
from datetime import datetime
import tempfile
import shutil
from openpyxl import load_workbook

# 设置日志配置
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('data_traversal.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# 定义全局变量
DATA_DIR = r"c:\Users\13067\Desktop\论文\非洲\demo\非洲数据"
OUTPUT_DIR = r"c:\Users\13067\Desktop\论文\非洲\demo\清洗数据"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "高级遍历数据.json")

# 确保输出目录存在
os.makedirs(OUTPUT_DIR, exist_ok=True)

# 存储所有提取的数据
all_articles = {}
processed_files = set()
error_files = []

# 生成唯一标识的函数
def generate_unique_key(article: Dict[str, Any]) -> str:
    url = article.get('url', '')
    title = article.get('title', '')
    publish_date = article.get('publish_date', '')
    content = article.get('content', '')[:100]  # 使用前100个字符作为内容摘要
    
    # 优先使用URL+标题+发布日期
    if url and title and publish_date:
        return f"{url}|{title}|{publish_date}"
    # 其次使用标题+发布日期+内容摘要
    elif title and publish_date:
        return f"{title}|{publish_date}|{content}"
    # 最后使用标题+内容摘要
    elif title:
        return f"{title}|{content}"
    # 最坏情况使用内容哈希
    else:
        import hashlib
        content_hash = hashlib.md5(content.encode('utf-8')).hexdigest()
        return content_hash

# 处理单个JSON对象
def process_json_object(obj: Any) -> List[Dict[str, Any]]:
    articles = []
    
    # 检查是否是包含results字段的对象
    if isinstance(obj, dict) and 'results' in obj:
        results = obj['results']
        if isinstance(results, list):
            articles.extend(results)
    # 检查是否直接是文章对象
    elif isinstance(obj, dict):
        articles.append(obj)
    # 检查是否是文章数组
    elif isinstance(obj, list):
        articles.extend(obj)
    
    return articles

# 处理JSON文件
def process_json_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing JSON file: {file_path}")
    articles = []
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read().strip()
        
        if not content:
            return {"file_path": file_path, "articles": [], "status": "empty"}
        
        # 尝试作为完整JSON解析
        try:
            data = json.loads(content)
            articles.extend(process_json_object(data))
        
        # 作为NDJSON处理（每行一个JSON对象）
        except json.JSONDecodeError:
            lines = content.split('\n')
            for line in lines:
                line = line.strip()
                if line:
                    try:
                        obj = json.loads(line)
                        articles.extend(process_json_object(obj))
                    except json.JSONDecodeError:
                        continue
        
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing JSON file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理GZ压缩文件
def process_gz_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing GZ file: {file_path}")
    articles = []
    
    try:
        with gzip.open(file_path, 'rt', encoding='utf-8') as f:
            content = f.read().strip()
        
        if not content:
            return {"file_path": file_path, "articles": [], "status": "empty"}
        
        # 尝试作为完整JSON解析
        try:
            data = json.loads(content)
            articles.extend(process_json_object(data))
        
        # 作为NDJSON处理
        except json.JSONDecodeError:
            lines = content.split('\n')
            for line in lines:
                line = line.strip()
                if line:
                    try:
                        obj = json.loads(line)
                        articles.extend(process_json_object(obj))
                    except json.JSONDecodeError:
                        continue
        
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing GZ file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理CSV文件
def process_csv_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing CSV file: {file_path}")
    articles = []
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                # 尝试将CSV行转换为文章格式
                article = {
                    "title": row.get('title', ''),
                    "url": row.get('url', ''),
                    "content": row.get('content', ''),
                    "source": row.get('source', ''),
                    "publish_date": row.get('publish_date', ''),
                    "keyword": row.get('keyword', ''),
                    "year": row.get('year', '')
                }
                # 只添加有内容的文章
                if article.get('content') or article.get('title'):
                    articles.append(article)
        
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing CSV file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理Excel文件
def process_excel_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing Excel file: {file_path}")
    articles = []
    
    try:
        workbook = load_workbook(file_path)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            # 读取表头
            headers = [cell.value for cell in sheet[1]]
            # 读取数据行
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(row):  # 跳过空行
                    row_dict = dict(zip(headers, row))
                    # 尝试将行转换为文章格式
                    article = {
                        "title": row_dict.get('title', ''),
                        "url": row_dict.get('url', ''),
                        "content": row_dict.get('content', ''),
                        "source": row_dict.get('source', ''),
                        "publish_date": row_dict.get('publish_date', ''),
                        "keyword": row_dict.get('keyword', ''),
                        "year": row_dict.get('year', '')
                    }
                    # 只添加有内容的文章
                    if article.get('content') or article.get('title'):
                        articles.append(article)
        
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing Excel file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理SQLite数据库文件
def process_db_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing DB file: {file_path}")
    articles = []
    
    try:
        # 连接数据库
        conn = sqlite3.connect(file_path)
        cursor = conn.cursor()
        
        # 获取所有表名
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
        tables = cursor.fetchall()
        
        for table in tables:
            table_name = table[0]
            logger.info(f"Processing table {table_name} in DB file {file_path}")
            
            # 获取表的列名
            cursor.execute(f"PRAGMA table_info({table_name});")
            columns = [col[1] for col in cursor.fetchall()]
            
            # 读取表中的所有数据
            cursor.execute(f"SELECT * FROM {table_name};")
            rows = cursor.fetchall()
            
            for row in rows:
                row_dict = dict(zip(columns, row))
                # 尝试将行转换为文章格式
                article = {
                    "title": row_dict.get('title', ''),
                    "url": row_dict.get('url', ''),
                    "content": row_dict.get('content', ''),
                    "source": row_dict.get('source', ''),
                    "publish_date": row_dict.get('publish_date', ''),
                    "keyword": row_dict.get('keyword', ''),
                    "year": row_dict.get('year', '')
                }
                # 只添加有内容的文章
                if article.get('content') or article.get('title'):
                    articles.append(article)
        
        conn.close()
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing DB file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理单个文件
def process_file(file_path: str) -> Dict[str, Any]:
    if file_path in processed_files:
        return {"file_path": file_path, "articles": [], "status": "already_processed"}
    
    processed_files.add(file_path)
    _, ext = os.path.splitext(file_path.lower())
    
    try:
        if ext in ['.json', '.backup']:
            return process_json_file(file_path)
        elif ext == '.gz':
            return process_gz_file(file_path)
        elif ext == '.csv':
            return process_csv_file(file_path)
        elif ext in ['.xlsx', '.xls']:
            return process_excel_file(file_path)
        elif ext == '.db':
            return process_db_file(file_path)
        else:
            logger.warning(f"Unsupported file type: {ext} for file {file_path}")
            return {"file_path": file_path, "articles": [], "status": "unsupported"}
    
    except Exception as e:
        logger.error(f"Unexpected error processing file {file_path}: {e}")
        error_files.append(file_path)
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 主处理函数
def main():
    global all_articles, processed_files, error_files
    
    logger.info("开始高级数据遍历...")
    start_time = datetime.now()
    
    # 获取所有文件
    logger.info("正在查找所有文件...")
    all_files = []
    for root, dirs, files in os.walk(DATA_DIR):
        for file in files:
            file_path = os.path.join(root, file)
            all_files.append(file_path)
    
    logger.info(f"找到 {len(all_files)} 个文件")
    
    # 使用多进程处理文件
    max_workers = os.cpu_count()
    logger.info(f"使用 {max_workers} 个进程处理文件")
    
    processed_count = 0
    with ProcessPoolExecutor(max_workers=max_workers) as executor:
        # 提交所有文件处理任务
        future_to_file = {executor.submit(process_file, file_path): file_path for file_path in all_files}
        
        # 处理完成的任务
        for future in as_completed(future_to_file):
            file_path = future_to_file[future]
            try:
                result = future.result()
                processed_count += 1
                
                # 处理提取的文章
                for article in result['articles']:
                    unique_key = generate_unique_key(article)
                    if unique_key not in all_articles:
                        all_articles[unique_key] = article
                
                # 每处理1000个文件打印一次进度
                if processed_count % 1000 == 0:
                    logger.info(f"已处理 {processed_count}/{len(all_files)} 个文件，当前唯一文章数: {len(all_articles)}")
            
            except Exception as e:
                logger.error(f"Error in future processing for file {file_path}: {e}")
                error_files.append(file_path)
    
    # 保存结果
    logger.info("保存处理结果...")
    unique_articles_list = list(all_articles.values())
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump(unique_articles_list, f, ensure_ascii=False, indent=2)
    
    # 保存处理状态
    processing_status = {
        "total_files": len(all_files),
        "processed_files": len(processed_files),
        "error_files": error_files,
        "total_articles": len(unique_articles_list),
        "processing_time": str(datetime.now() - start_time)
    }
    with open(os.path.join(OUTPUT_DIR, "处理状态.json"), 'w', encoding='utf-8') as f:
        json.dump(processing_status, f, ensure_ascii=False, indent=2)
    
    logger.info("高级数据遍历完成！")
    logger.info(f"总文件数: {len(all_files)}")
    logger.info(f"成功处理文件数: {len(processed_files) - len(error_files)}")
    logger.info(f"失败文件数: {len(error_files)}")
    logger.info(f"唯一文章数: {len(unique_articles_list)}")
    logger.info(f"输出文件: {OUTPUT_FILE}")
    logger.info(f"处理时间: {datetime.now() - start_time}")

if __name__ == "__main__":
    main()