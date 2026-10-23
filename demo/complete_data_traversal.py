import os
import json
import gzip
import sqlite3
import logging
import hashlib
import time
from typing import Dict, List, Any, Optional
from datetime import datetime
import tempfile
import shutil

# 设置日志配置
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('complete_data_traversal.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# 定义全局变量
DATA_DIR = r"c:\Users\13067\Desktop\论文\非洲\demo\非洲数据\非洲数据"
OUTPUT_DIR = r"c:\Users\13067\Desktop\论文\非洲\demo\清洗数据"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "原始数据.json")
CHECKPOINT_FILE = os.path.join(OUTPUT_DIR, "遍历 checkpoint.json")

# 确保输出目录存在
os.makedirs(OUTPUT_DIR, exist_ok=True)

# 生成唯一标识的函数
def generate_unique_key(article: Any) -> str:
    # 确保article是字典类型
    if not isinstance(article, dict):
        # 对于非字典类型，使用其字符串表示作为唯一标识
        return str(article)
    
    url = article.get('url', '')
    title = article.get('title', '')
    publish_date = article.get('publish_date', '')
    content = article.get('content', '')[:200]  # 使用前200个字符作为内容摘要
    
    # 复合键去重：URL + 标题 + 发布日期 + 内容摘要
    key_parts = [url, title, publish_date, content]
    return "|" .join([part for part in key_parts if part])

# 处理单个JSON对象
def process_json_object(obj: Any) -> List[Dict[str, Any]]:
    articles = []
    
    # 检查是否是包含results字段的对象
    if isinstance(obj, dict) and 'results' in obj:
        results = obj['results']
        if isinstance(results, list):
            articles.extend(results)
        elif isinstance(results, dict):
            # 处理可能嵌套的results字典
            articles.extend(process_json_object(results))
    # 检查是否包含data字段的对象
    elif isinstance(obj, dict) and 'data' in obj:
        data = obj['data']
        if isinstance(data, list):
            articles.extend(data)
        elif isinstance(data, dict):
            articles.extend(process_json_object(data))
    # 检查是否包含articles字段的对象
    elif isinstance(obj, dict) and 'articles' in obj:
        articles_field = obj['articles']
        if isinstance(articles_field, list):
            articles.extend(articles_field)
        elif isinstance(articles_field, dict):
            articles.extend(process_json_object(articles_field))
    # 检查是否直接是文章对象
    elif isinstance(obj, dict):
        articles.append(obj)
    # 检查是否是文章数组
    elif isinstance(obj, list):
        articles.extend(obj)
    
    return articles

# 处理JSON文件
def process_json_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing JSON file: {file_path}")
    articles = []
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read().strip()
        
        if not content:
            return {"file_path": file_path, "articles": [], "status": "empty"}
        
        # 尝试作为完整JSON解析
        try:
            data = json.loads(content)
            articles.extend(process_json_object(data))
        
        # 作为NDJSON处理（每行一个JSON对象）
        except json.JSONDecodeError:
            lines = content.split('\n')
            for line in lines:
                line = line.strip()
                if line:
                    try:
                        obj = json.loads(line)
                        articles.extend(process_json_object(obj))
                    except json.JSONDecodeError:
                        continue
        
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing JSON file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理GZ压缩文件
def process_gz_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing GZ file: {file_path}")
    articles = []
    
    try:
        with gzip.open(file_path, 'rt', encoding='utf-8') as f:
            content = f.read().strip()
        
        if not content:
            return {"file_path": file_path, "articles": [], "status": "empty"}
        
        # 尝试作为完整JSON解析
        try:
            data = json.loads(content)
            articles.extend(process_json_object(data))
        
        # 作为NDJSON处理
        except json.JSONDecodeError:
            lines = content.split('\n')
            for line in lines:
                line = line.strip()
                if line:
                    try:
                        obj = json.loads(line)
                        articles.extend(process_json_object(obj))
                    except json.JSONDecodeError:
                        continue
        
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing GZ file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理CSV文件
def process_csv_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing CSV file: {file_path}")
    import csv
    articles = []
    
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                # 尝试将CSV行转换为文章格式
                article = {
                    "title": row.get('title', ''),
                    "url": row.get('url', ''),
                    "content": row.get('content', ''),
                    "source": row.get('source', ''),
                    "publish_date": row.get('publish_date', ''),
                    "keyword": row.get('keyword', ''),
                    "year": row.get('year', '')
                }
                # 只添加有内容的文章
                if article.get('content') or article.get('title'):
                    articles.append(article)
        
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing CSV file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理Excel文件
def process_excel_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing Excel file: {file_path}")
    articles = []
    
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(file_path)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            # 读取表头
            headers = [cell.value for cell in sheet[1]]
            # 读取数据行
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(row):  # 跳过空行
                    row_dict = dict(zip(headers, row))
                    # 尝试将行转换为文章格式
                    article = {
                        "title": row_dict.get('title', ''),
                        "url": row_dict.get('url', ''),
                        "content": row_dict.get('content', ''),
                        "source": row_dict.get('source', ''),
                        "publish_date": row_dict.get('publish_date', ''),
                        "keyword": row_dict.get('keyword', ''),
                        "year": row_dict.get('year', '')
                    }
                    # 只添加有内容的文章
                    if article.get('content') or article.get('title'):
                        articles.append(article)
        
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing Excel file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理SQLite数据库文件
def process_db_file(file_path: str) -> Dict[str, Any]:
    logger.info(f"Processing DB file: {file_path}")
    articles = []
    
    try:
        # 连接数据库
        conn = sqlite3.connect(file_path)
        cursor = conn.cursor()
        
        # 获取所有表名
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
        tables = cursor.fetchall()
        
        for table in tables:
            table_name = table[0]
            logger.info(f"Processing table {table_name} in DB file {file_path}")
            
            # 获取表的列名
            cursor.execute(f"PRAGMA table_info({table_name});")
            columns = [col[1] for col in cursor.fetchall()]
            
            # 读取表中的所有数据
            cursor.execute(f"SELECT * FROM {table_name};")
            rows = cursor.fetchall()
            
            for row in rows:
                row_dict = dict(zip(columns, row))
                # 尝试将行转换为文章格式
                article = {
                    "title": row_dict.get('title', ''),
                    "url": row_dict.get('url', ''),
                    "content": row_dict.get('content', ''),
                    "source": row_dict.get('source', ''),
                    "publish_date": row_dict.get('publish_date', ''),
                    "keyword": row_dict.get('keyword', ''),
                    "year": row_dict.get('year', '')
                }
                # 只添加有内容的文章
                if article.get('content') or article.get('title'):
                    articles.append(article)
        
        conn.close()
        return {"file_path": file_path, "articles": articles, "status": "success"}
    
    except Exception as e:
        logger.error(f"Error processing DB file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 处理单个文件
def process_file(file_path: str) -> Dict[str, Any]:
    file_name = os.path.basename(file_path).lower()
    _, ext = os.path.splitext(file_path.lower())
    
    try:
        # 支持多种备份文件类型
        if ext in ['.json', '.backup', '.current_backup'] or file_name.endswith('.json.backup') or file_name.endswith('.json.current_backup'):
            return process_json_file(file_path)
        elif ext == '.gz':
            return process_gz_file(file_path)
        elif ext == '.csv':
            return process_csv_file(file_path)
        elif ext in ['.xlsx', '.xls']:
            return process_excel_file(file_path)
        elif ext == '.db':
            return process_db_file(file_path)
        else:
            logger.warning(f"Unsupported file type: {ext} for file {file_path}")
            return {"file_path": file_path, "articles": [], "status": "unsupported"}
    
    except Exception as e:
        logger.error(f"Unexpected error processing file {file_path}: {e}")
        return {"file_path": file_path, "articles": [], "status": "error", "error": str(e)}

# 加载断点
def load_checkpoint() -> Dict[str, Any]:
    if os.path.exists(CHECKPOINT_FILE):
        try:
            with open(CHECKPOINT_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            logger.error(f"Error loading checkpoint: {e}")
    return {"processed_files": [], "all_articles": {}}

# 保存断点
def save_checkpoint(processed_files: List[str], all_articles: Dict[str, Any]) -> None:
    checkpoint = {
        "processed_files": processed_files,
        "all_articles": all_articles,
        "timestamp": datetime.now().isoformat()
    }
    try:
        with open(CHECKPOINT_FILE, 'w', encoding='utf-8') as f:
            json.dump(checkpoint, f, ensure_ascii=False, indent=2)
        logger.info(f"Checkpoint saved, processed {len(processed_files)} files, {len(all_articles)} unique articles")
    except Exception as e:
        logger.error(f"Error saving checkpoint: {e}")

# 主处理函数
def main():
    logger.info("开始完整数据遍历...")
    start_time = datetime.now()
    
    # 重新开始处理，不使用断点
    processed_files = set()
    all_articles = {}
    
    logger.info(f"从头开始处理，当前唯一文章数: {len(all_articles)}")
    
    # 获取所有文件
    logger.info("正在查找所有文件...")
    all_files = []
    for root, dirs, files in os.walk(DATA_DIR):
        for file in files:
            file_path = os.path.join(root, file)
            all_files.append(file_path)
    
    logger.info(f"找到 {len(all_files)} 个文件")
    
    # 处理所有文件
    files_to_process = all_files
    logger.info(f"需要处理 {len(files_to_process)} 个文件")
    
    # 处理文件
    processed_count = 0
    for file_path in files_to_process:
        result = process_file(file_path)
        processed_files.add(file_path)
        processed_count += 1
        
        # 处理提取的文章
        for article in result['articles']:
            unique_key = generate_unique_key(article)
            if unique_key not in all_articles:
                all_articles[unique_key] = article
        
        # 每处理200个文件保存一次断点
        if processed_count % 200 == 0:
            save_checkpoint(list(processed_files), all_articles)
        
        # 每处理1000个文件打印一次进度
        if processed_count % 1000 == 0:
            logger.info(f"已处理 {processed_count}/{len(files_to_process)} 个文件，当前唯一文章数: {len(all_articles)}")
    
    # 最终保存
    save_checkpoint(list(processed_files), all_articles)
    
    # 保存结果
    logger.info("保存最终结果...")
    unique_articles_list = list(all_articles.values())
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump(unique_articles_list, f, ensure_ascii=False, indent=2)
    
    # 生成处理报告
    processing_time = datetime.now() - start_time
    report = {
        "total_files": len(all_files),
        "processed_files": len(processed_files),
        "total_articles": len(unique_articles_list),
        "processing_time": str(processing_time),
        "start_time": start_time.isoformat(),
        "end_time": datetime.now().isoformat(),
        "output_file": OUTPUT_FILE
    }
    
    report_file = os.path.join(OUTPUT_DIR, "遍历处理报告.json")
    with open(report_file, 'w', encoding='utf-8') as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    
    logger.info("完整数据遍历完成！")
    logger.info(f"总文件数: {report['total_files']}")
    logger.info(f"已处理文件数: {report['processed_files']}")
    logger.info(f"唯一文章数: {report['total_articles']}")
    logger.info(f"处理时间: {report['processing_time']}")
    logger.info(f"输出文件: {report['output_file']}")
    logger.info(f"处理报告: {report_file}")

if __name__ == "__main__":
    main()