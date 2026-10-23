# -*- coding: utf-8 -*-
"""
Generate Sample Profile Table in Nature/Science Journal Style
Generates LaTeX (booktabs), Word (.docx), and CSV format tables
"""

import pandas as pd
import numpy as np
from pathlib import Path
import warnings
warnings.filterwarnings('ignore')

DATA_PATH = Path(r"c:\Users\13067\Desktop\论文\非洲\demo\清洗数据\enhanced_data_with_complete_features.csv")
OUTPUT_DIR = Path(r"c:\Users\13067\Desktop\论文\非洲")

REGION_MAPPING = {
    '北非': 'North Africa',
    '东非': 'East Africa',
    '西非': 'West Africa',
    '中非': 'Central Africa',
    '南非': 'Southern Africa'
}

CATEGORY_MAPPING = {
    'Business - Economy': 'Business & Economy',
    'World - International': 'World & International',
    'Politics': 'Politics',
    'Sports': 'Sports',
    'Sports - World': 'Sports',
    'Culture - Arts': 'Culture & Arts',
    'Technology - Science': 'Technology & Science',
    'Health': 'Health',
    'Opinion': 'Opinion',
    'Opinion - World': 'Opinion',
    'Opinion - Opinion': 'Opinion',
    'Weekly - World': 'Weekly',
    'Weekly - Heritage': 'Weekly',
    'Weekly - Egypt': 'Weekly',
    'Egypt - Society': 'Egypt Society',
    'Business - Energy': 'Business & Economy',
    'Other': 'Other',
    'Unknown': 'Unknown'
}

SENTIMENT_MAPPING = {
    'positive': 'Positive',
    'negative': 'Negative',
    'neutral': 'Neutral'
}


def load_data():
    df = pd.read_csv(DATA_PATH, low_memory=False)
    df['category'] = df['category'].fillna('Unknown')
    df['category'] = df['category'].astype(str).str.replace(r'\s+', ' ', regex=True).str.strip()
    df['category'] = df['category'].replace('nan', 'Unknown')
    return df


def calc_categorical_stats(df, col, mapping=None):
    counts = df[col].value_counts()
    total = len(df)
    
    if mapping:
        mapped_counts = {}
        for val, count in counts.items():
            if val in mapping:
                label = mapping[val]
            else:
                label = str(val)
            if label in mapped_counts:
                mapped_counts[label] += count
            else:
                mapped_counts[label] = count
        results = [{'category': k, 'n': v, 'percentage': v/total*100} 
                   for k, v in mapped_counts.items()]
    else:
        results = [{'category': str(val), 'n': count, 'percentage': count/total*100}
                   for val, count in counts.items()]
    
    return sorted(results, key=lambda x: -x['n'])


def calc_continuous_stats(df, col):
    data = df[col].dropna()
    mean_val = data.mean()
    std_val = data.std()
    return mean_val, std_val


def escape_latex(text):
    text = str(text)
    text = text.replace('\\', '\\textbackslash{}')
    text = text.replace('&', '\\&')
    text = text.replace('%', '\\%')
    text = text.replace('$', '\\$')
    text = text.replace('#', '\\#')
    text = text.replace('_', '\\_')
    text = text.replace('{', '\\{')
    text = text.replace('}', '\\}')
    return text


def generate_latex_table(sample_size, year_stats, region_stats, category_stats, 
                         sentiment_stats, continuous_vars):
    latex_content = r"""\documentclass[11pt]{article}
\usepackage[utf8]{inputenc}
\usepackage{booktabs}
\usepackage{geometry}
\usepackage{array}
\geometry{a4paper, margin=1in}

\begin{document}

\begin{table}[htbp]
\centering
\caption{Sample Profile}
\label{tab:sample_profile}
\begin{tabular}{llr}
\toprule
\textbf{Variable} & \textbf{Category} & \textbf{n (\%)} \\
\midrule
\multicolumn{3}{l}{\textit{Total Sample Size}} \\
& & \textbf{""" + f"{sample_size:,}" + r"""} \\
\midrule
\multicolumn{3}{l}{\textit{Publication Year}} \\
"""
    
    for stat in year_stats:
        cat = escape_latex(stat['category'])
        latex_content += f"& {cat} & {stat['n']:,} ({stat['percentage']:.1f}\\%) \\\\\n"
    
    latex_content += r"""\midrule
\multicolumn{3}{l}{\textit{Source Region}} \\
"""
    
    for stat in region_stats:
        cat = escape_latex(stat['category'])
        latex_content += f"& {cat} & {stat['n']:,} ({stat['percentage']:.1f}\\%) \\\\\n"
    
    latex_content += r"""\midrule
\multicolumn{3}{l}{\textit{News Category}} \\
"""
    
    for stat in category_stats:
        cat = escape_latex(stat['category'])
        latex_content += f"& {cat} & {stat['n']:,} ({stat['percentage']:.1f}\\%) \\\\\n"
    
    latex_content += r"""\midrule
\multicolumn{3}{l}{\textit{Sentiment Distribution}} \\
"""
    
    for stat in sentiment_stats:
        cat = escape_latex(stat['category'])
        latex_content += f"& {cat} & {stat['n']:,} ({stat['percentage']:.1f}\\%) \\\\\n"
    
    latex_content += r"""\midrule
\multicolumn{3}{l}{\textit{Continuous Variables (Mean $\pm$ SD)}} \\
"""
    
    for var_name, mean_val, std_val in continuous_vars:
        var = escape_latex(var_name)
        latex_content += f"& {var} & {mean_val:.2f} $\\pm$ {std_val:.2f} \\\\\n"
    
    latex_content += r"""\bottomrule
\end{tabular}
\end{table}

\end{document}
"""
    
    return latex_content


def generate_word_table(sample_size, year_stats, region_stats, category_stats,
                        sentiment_stats, continuous_vars, output_path):
    try:
        from docx import Document
        from docx.shared import Pt, Inches, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        print("python-docx not installed. Installing...")
        import subprocess
        subprocess.check_call(['pip', 'install', 'python-docx'])
        from docx import Document
        from docx.shared import Pt, Inches, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    
    doc = Document()
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    
    title = doc.add_paragraph()
    title_run = title.add_run('Table 1. Sample Profile')
    title_run.bold = True
    title_run.font.name = 'Times New Roman'
    title_run.font.size = Pt(11)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    total_rows = (1 + 1 + len(year_stats) + 1 + len(region_stats) + 1 + 
                  len(category_stats) + 1 + len(sentiment_stats) + 1 + 
                  len(continuous_vars) + 1)
    
    table = doc.add_table(rows=total_rows, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    def set_cell_font(cell, text, bold=False, italic=False):
        cell.text = text
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                run.bold = bold
                run.italic = italic
    
    def add_top_border(row_idx):
        row = table.rows[row_idx]
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            top = OxmlElement('w:top')
            top.set(qn('w:val'), 'single')
            top.set(qn('w:sz'), '12')
            tcBorders.append(top)
            tcPr.append(tcBorders)
    
    def add_bottom_border(row_idx, thick=True):
        row = table.rows[row_idx]
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '12' if thick else '6')
            tcBorders.append(bottom)
            tcPr.append(tcBorders)
    
    row_idx = 0
    
    set_cell_font(table.rows[row_idx].cells[0], 'Variable', bold=True)
    set_cell_font(table.rows[row_idx].cells[1], 'Category', bold=True)
    set_cell_font(table.rows[row_idx].cells[2], 'n (%)', bold=True)
    add_top_border(row_idx)
    add_bottom_border(row_idx)
    row_idx += 1
    
    set_cell_font(table.rows[row_idx].cells[0], 'Total Sample Size', italic=True)
    set_cell_font(table.rows[row_idx].cells[1], '')
    set_cell_font(table.rows[row_idx].cells[2], f'{sample_size:,}', bold=True)
    row_idx += 1
    
    set_cell_font(table.rows[row_idx].cells[0], 'Publication Year', italic=True)
    add_bottom_border(row_idx, thick=False)
    row_idx += 1
    
    for stat in year_stats:
        set_cell_font(table.rows[row_idx].cells[0], '')
        set_cell_font(table.rows[row_idx].cells[1], stat['category'])
        set_cell_font(table.rows[row_idx].cells[2], f"{stat['n']:,} ({stat['percentage']:.1f}%)")
        row_idx += 1
    
    set_cell_font(table.rows[row_idx].cells[0], 'Source Region', italic=True)
    add_bottom_border(row_idx, thick=False)
    row_idx += 1
    
    for stat in region_stats:
        set_cell_font(table.rows[row_idx].cells[0], '')
        set_cell_font(table.rows[row_idx].cells[1], stat['category'])
        set_cell_font(table.rows[row_idx].cells[2], f"{stat['n']:,} ({stat['percentage']:.1f}%)")
        row_idx += 1
    
    set_cell_font(table.rows[row_idx].cells[0], 'News Category', italic=True)
    add_bottom_border(row_idx, thick=False)
    row_idx += 1
    
    for stat in category_stats:
        set_cell_font(table.rows[row_idx].cells[0], '')
        set_cell_font(table.rows[row_idx].cells[1], stat['category'])
        set_cell_font(table.rows[row_idx].cells[2], f"{stat['n']:,} ({stat['percentage']:.1f}%)")
        row_idx += 1
    
    set_cell_font(table.rows[row_idx].cells[0], 'Sentiment Distribution', italic=True)
    add_bottom_border(row_idx, thick=False)
    row_idx += 1
    
    for stat in sentiment_stats:
        set_cell_font(table.rows[row_idx].cells[0], '')
        set_cell_font(table.rows[row_idx].cells[1], stat['category'])
        set_cell_font(table.rows[row_idx].cells[2], f"{stat['n']:,} ({stat['percentage']:.1f}%)")
        row_idx += 1
    
    set_cell_font(table.rows[row_idx].cells[0], 'Continuous Variables', italic=True)
    set_cell_font(table.rows[row_idx].cells[1], '(Mean ± SD)')
    add_bottom_border(row_idx, thick=False)
    row_idx += 1
    
    for var_name, mean_val, std_val in continuous_vars:
        set_cell_font(table.rows[row_idx].cells[0], '')
        set_cell_font(table.rows[row_idx].cells[1], var_name)
        set_cell_font(table.rows[row_idx].cells[2], f"{mean_val:.2f} ± {std_val:.2f}")
        row_idx += 1
    
    add_bottom_border(row_idx - 1, thick=True)
    
    doc.save(output_path)
    print(f"Word document saved: {output_path}")


def generate_csv_table(sample_size, year_stats, region_stats, category_stats,
                       sentiment_stats, continuous_vars, output_path):
    rows = []
    
    rows.append(['Variable', 'Category', 'n', '%'])
    rows.append(['Total Sample Size', '', sample_size, ''])
    rows.append(['', '', '', ''])
    
    rows.append(['Publication Year', '', '', ''])
    for stat in year_stats:
        rows.append(['', stat['category'], stat['n'], f"{stat['percentage']:.1f}"])
    rows.append(['', '', '', ''])
    
    rows.append(['Source Region', '', '', ''])
    for stat in region_stats:
        rows.append(['', stat['category'], stat['n'], f"{stat['percentage']:.1f}"])
    rows.append(['', '', '', ''])
    
    rows.append(['News Category', '', '', ''])
    for stat in category_stats:
        rows.append(['', stat['category'], stat['n'], f"{stat['percentage']:.1f}"])
    rows.append(['', '', '', ''])
    
    rows.append(['Sentiment Distribution', '', '', ''])
    for stat in sentiment_stats:
        rows.append(['', stat['category'], stat['n'], f"{stat['percentage']:.1f}"])
    rows.append(['', '', '', ''])
    
    rows.append(['Continuous Variables', '', 'Mean', 'SD'])
    for var_name, mean_val, std_val in continuous_vars:
        rows.append(['', var_name, f"{mean_val:.2f}", f"{std_val:.2f}"])
    
    df_out = pd.DataFrame(rows)
    df_out.to_csv(output_path, index=False, header=False)
    print(f"CSV file saved: {output_path}")


def main():
    print("Loading data...")
    df = load_data()
    sample_size = len(df)
    print(f"Total samples: {sample_size:,}")
    
    print("\nCalculating categorical statistics...")
    year_stats = calc_categorical_stats(df, 'year')
    
    region_stats = calc_categorical_stats(df, 'source_region_5', REGION_MAPPING)
    
    category_stats = calc_categorical_stats(df, 'category', CATEGORY_MAPPING)
    
    sentiment_stats = calc_categorical_stats(df, 'sentiment_label', SENTIMENT_MAPPING)
    
    print("Calculating continuous variable statistics...")
    continuous_vars = [
        ('Normalized Sentiment', *calc_continuous_stats(df, 'normalized_sentiment')),
        ('China-Africa Community Recognition', *calc_continuous_stats(df, 'china_africa_community_recognition')),
        ('Development Demand (PCA)', *calc_continuous_stats(df, 'development_demand_pca')),
        ('Cooperation Depth Index', *calc_continuous_stats(df, 'cooperation_depth_index')),
        ('Long-term Word Count', *calc_continuous_stats(df, 'long_term_word_count'))
    ]
    
    print("\n" + "="*60)
    print("SAMPLE PROFILE SUMMARY")
    print("="*60)
    print(f"\nTotal Sample Size: {sample_size:,}")
    
    print("\n--- Publication Year Distribution ---")
    for stat in year_stats[:10]:
        print(f"  {stat['category']}: {stat['n']:,} ({stat['percentage']:.1f}%)")
    
    print("\n--- Source Region Distribution ---")
    for stat in region_stats:
        print(f"  {stat['category']}: {stat['n']:,} ({stat['percentage']:.1f}%)")
    
    print("\n--- News Category Distribution ---")
    for stat in category_stats[:10]:
        print(f"  {stat['category']}: {stat['n']:,} ({stat['percentage']:.1f}%)")
    
    print("\n--- Sentiment Distribution ---")
    for stat in sentiment_stats:
        print(f"  {stat['category']}: {stat['n']:,} ({stat['percentage']:.1f}%)")
    
    print("\n--- Continuous Variables (Mean ± SD) ---")
    for var_name, mean_val, std_val in continuous_vars:
        print(f"  {var_name}: {mean_val:.2f} ± {std_val:.2f}")
    
    print("\n" + "="*60)
    print("GENERATING OUTPUT FILES")
    print("="*60)
    
    latex_content = generate_latex_table(sample_size, year_stats, region_stats,
                                         category_stats, sentiment_stats, continuous_vars)
    latex_path = OUTPUT_DIR / "sample_profile_table.tex"
    with open(latex_path, 'w', encoding='utf-8') as f:
        f.write(latex_content)
    print(f"LaTeX file saved: {latex_path}")
    
    word_path = OUTPUT_DIR / "sample_profile_table.docx"
    generate_word_table(sample_size, year_stats, region_stats, category_stats,
                        sentiment_stats, continuous_vars, word_path)
    
    csv_path = OUTPUT_DIR / "sample_profile_table.csv"
    generate_csv_table(sample_size, year_stats, region_stats, category_stats,
                       sentiment_stats, continuous_vars, csv_path)
    
    print("\n" + "="*60)
    print("ALL FILES GENERATED SUCCESSFULLY!")
    print("="*60)
    print(f"\nOutput files:")
    print(f"  1. LaTeX: {latex_path}")
    print(f"  2. Word:  {word_path}")
    print(f"  3. CSV:   {csv_path}")


if __name__ == "__main__":
    main()
